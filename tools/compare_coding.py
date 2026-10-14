"""Compare a second coder's template with the reference coding and report agreement.

Usage:
  python3 tools/compare_coding.py coding/ARWC_Coding_Round1_<coder>.xlsx data/ARWC_Case_Inventory.xlsx coding/Agreement_Round1_<coder>.xlsx

Scored fields are compared after normalising to the codebook vocabulary. Free-text fields are
listed side by side as "Review" and left to discussion. Percentage agreement and Cohen's kappa
for the checkbox items (categories, factors, triggers) are Excel formulas in the report.
"""
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule

F = "Arial"
HDR = PatternFill("solid", start_color="2F4A8A")

SENTENCE = ["Death (immediate)", "Death with reprieve", "Life imprisonment", "Fixed term", "Other"]
AGE = ["Minor", "Adult", "Mixed", "Unknown"]
REL = ["Stranger", "Acquaintance", "Family member", "Other", "Unknown"]
FIN = ["Ordinary appeals completed", "Ordinary appeal period expired", "Judgment otherwise entered into legal effect"]
ACTION = ["Retrial acquittal", "Post-finality appellate acquittal", "Conviction vacated or annulled",
          "Extraordinary appeal decision quashing a final judgment", "Prosecutorial admission of error",
          "State compensation decision", "Administrative rectification", "Pardon", "Commutation", "Other"]
BASIS = ["Insufficiency of the original evidence", "New evidence of innocence", "Procedural or legal error", "Other", "Not stated"]
EVID = ["confession", "eyewitness", "forensic", "documentary", "informant", "circumstantial", "other"]


def s(v):
    return "" if v is None else str(v).strip()


def cv(v, vocab):
    """Map a value to the vocabulary entry it starts with (case-insensitive)."""
    low = s(v).lower()
    for item in sorted(vocab, key=len, reverse=True):
        if low.startswith(item.lower()):
            return item
    return s(v)


def yn(v):
    low = s(v).lower()
    if low.startswith(("yes", "marked", "☒")):
        return "Yes"
    if low.startswith(("no", "not marked", "☐")):
        return "No"
    if low.startswith("unknown"):
        return "Unknown"
    return s(v)


def date(v):
    t = s(v)
    m = re.search(r"Year:\s*(\d{4})(?:\s*\|\s*Month:\s*(\d{1,2}))?(?:\s*\|\s*Day:\s*(\d{1,2}))?", t)
    if m:
        return "-".join(x.zfill(2) if i else x for i, x in enumerate(m.groups()) if x)
    m = re.search(r"(\d{4})(?:[-/.](\d{1,2}))?(?:[-/.](\d{1,2}))?", t)
    return "-".join(x.zfill(2) if i else x for i, x in enumerate(m.groups()) if x) if m else t


def first_int(v):
    m = re.search(r"\d+", s(v))
    return int(m.group(0)) if m else None


def first_float(v):
    m = re.search(r"\d+(?:\.\d+)?", s(v))
    return float(m.group(0)) if m else None


def tokens(v, mapper=None):
    out = set()
    for part in re.split(r"[;,]", s(v)):
        part = re.sub(r"\(.*?\)", "", part).strip().lower()
        if not part:
            continue
        if mapper:
            part = next((e for e in mapper if part.startswith(e)), part)
        out.add(part)
    return out


def compare(kind, coder, ref):
    """Return (result, coder_norm, ref_norm). Result: Agree / Partly / Disagree / Review / Not coded."""
    if kind == "review":
        return "Review", s(coder), s(ref)
    if not s(coder):
        return "Not coded", "", s(ref)
    if kind == "date":
        c, r = date(coder), date(ref)
        if c == r:
            return "Agree", c, r
        return ("Partly" if c[:4] == r[:4] and c[:4] else "Disagree"), c, r
    if kind in ("int", "float"):
        c, r = (first_int, first_float)[kind == "float"](coder), (first_int, first_float)[kind == "float"](ref)
        if c is None or r is None:
            return "Review", s(coder), s(ref)
        ok = abs(c - r) <= (0.15 if kind == "float" else 0)
        return ("Agree" if ok else "Disagree"), c, r
    if kind in ("set", "evidence"):
        c, r = tokens(coder, EVID if kind == "evidence" else None), tokens(ref, EVID if kind == "evidence" else None)
        res = "Agree" if c == r else ("Partly" if c & r else "Disagree")
        return res, "; ".join(sorted(c)), "; ".join(sorted(r))
    if kind == "yn":
        c, r = yn(coder), yn(ref)
    elif kind == "upper":
        c, r = s(coder).upper(), s(ref).upper()
    elif kind == "cat":
        c, r = (s(coder)[:1].upper() if s(coder)[:1].upper() in "ABC" else "None"), (s(ref)[:1].upper() if s(ref)[:1].upper() in "ABC" else "None")
    else:
        vocab = {"sentence": SENTENCE, "age": AGE, "rel": REL, "fin": FIN, "action": ACTION, "basis": BASIS}[kind]
        c, r = cv(coder, vocab), cv(ref, vocab)
    return ("Agree" if c.lower() == r.lower() else "Disagree"), c, r


# key in the template -> comparison kind (keys match the reference sheet's column names)
CASE_FIELDS = {
    "Case Name (Original Language)": "review", "Original Language": "review",
    "Offense Category": "set", "Offense as Charged (Original Language)": "review", "Offense as Charged (English Gloss)": "review",
    "Sentence": "sentence", "Court of Original Trial": "review", "Evidence Relied Upon at Trial": "evidence",
    "Number of Victims": "int", "Victim Age Group": "age", "Victim–Defendant Relationship": "rel",
    "Alleged Victim Later Reappeared": "yn", "Actual Perpetrator Identified": "yn",
    "Subsequent Victims Attributed to Actual Perpetrator": "yn",
    "Finality Basis": "fin", "Date of Finality": "date", "Finality Note": "review",
    "Compensation: Compensation Awarded": "yn", "Compensation: Amount / Currency / Year": "review",
    "Compensation: Decision Acknowledges Wrongful Conviction": "yn",
    "Criminal offence: status": "upper", "Finality: status": "upper", "Official post-conviction action: status": "upper",
    "Documentation: status": "upper", "Inclusion determination: status": "upper",
}
IND_FIELDS = {"Individual Name (Original Language)": "review", "Gender": "yn_gender", "Age at Conviction": "int", "Executed": "yn",
              "Date of Execution": "date", "Years Lost": "float", "Years Lost: Approximation Flag": "yn", "Posthumous Case": "yn"}
ACT_FIELDS = {"Action Type": "action", "Category": "cat", "Date of Action": "date", "Basis Stated in the Decision": "basis",
              "Document Type": "review", "Issuing Authority": "review", "Procedural Route": "review"}


def read_coder(path):
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if not re.match(r"^(ARWC|EXC)-", ws.title):
            continue
        vals = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                vals[s(row[0])] = (row[3], s(row[4]), s(row[2]), s(row[1]))
        out[ws.title] = vals
    return out


def read_reference(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Coded cases"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    cases = {r[hdr.index("Case ID")]: dict(zip(hdr, r)) for r in ws.iter_rows(min_row=3, values_only=True)}
    acts, inds = {}, {}
    ah = [c.value for c in next(wb["Official actions"].iter_rows(max_row=1))]
    for r in wb["Official actions"].iter_rows(min_row=2, values_only=True):
        acts.setdefault(r[0], []).append(dict(zip(ah, r)))
    ih = [c.value for c in next(wb["Individuals"].iter_rows(max_row=1))]
    for r in wb["Individuals"].iter_rows(min_row=2, values_only=True):
        inds.setdefault(r[0], []).append(dict(zip(ih, r)))
    return cases, acts, inds


def main(coder_path, ref_path, out_path):
    coder = read_coder(coder_path)
    ref, ref_acts, ref_inds = read_reference(ref_path)
    rows, binary, logs = [], [], []
    for cid, vals in coder.items():
        r = ref.get(cid)
        if r is None:
            print(f"warning: {cid} has no reference coding; skipped")
            continue
        get = lambda k: vals.get(k, ("", "", k, ""))
        # categories, factors, triggers: checkbox items, also scored for kappa
        for key in [k for k in vals if k in ("Category A", "Category B", "Category C") or k.endswith(": marked")]:
            cval, note, label, sec = get(key)
            res, cn, rn = compare("yn", cval, r.get(key))
            rows.append([cid, sec, label, cn, note, rn, res])
            if cn in ("Yes", "No") and rn in ("Yes", "No"):
                group = "Category" if key.startswith("Category") else ("Factor" if re.match(r"^\d", key) else "Trigger")
                binary.append([cid, group, label, int(cn == "Yes"), int(rn == "Yes")])
        for key, kind in CASE_FIELDS.items():
            if key not in vals:
                continue
            cval, note, label, sec = get(key)
            res, cn, rn = compare(kind, cval, r.get(key))
            rows.append([cid, sec, label, cn, note, rn, res])
        for i, ri in enumerate(ref_inds.get(cid, []), 1):
            for f, kind in IND_FIELDS.items():
                key = f"B{i}|{f}"
                if key not in vals:
                    continue
                cval, note, label, sec = get(key)
                if kind == "yn_gender":
                    res, cn, rn = ("Agree" if s(cval).lower() == s(ri.get(f)).lower().split(" ")[0] else ("Not coded" if not s(cval) else "Disagree"), s(cval), s(ri.get(f)))
                else:
                    res, cn, rn = compare(kind, cval, ri.get(f))
                rows.append([cid, sec, label, cn, note, rn, res])
        # official actions: pair by category, then compare type, date and basis
        coder_acts = []
        for a in range(1, 10):
            if f"Act{a}|Action Type" not in vals:
                break
            if s(get(f"Act{a}|Action Type")[0]) or s(get(f"Act{a}|Category")[0]):
                coder_acts.append({f: get(f"Act{a}|{f}")[0] for f in ACT_FIELDS} | {"_note": get(f"Act{a}|Basis Stated in the Decision")[1], "_n": a})
        refs = ref_acts.get(cid, [])
        res = "Agree" if len(coder_acts) == len(refs) else "Disagree"
        rows.append([cid, "E. Official actions", "Number of official actions", len(coder_acts), "", len(refs), res])
        used = set()
        for ra in refs:
            rcat = s(ra.get("Category"))[:1].upper() or "None"
            match = next((ca for ca in coder_acts if ca["_n"] not in used and s(ca["Category"])[:1].upper() == rcat[:1]), None)
            label = s(ra.get("Action")).split(" — ")[0] + f" (reference category {rcat})"
            if not match:
                rows.append([cid, "E. Official actions", label, "", "", s(ra.get("Action Type")), "Disagree (no matching action)"])
                continue
            used.add(match["_n"])
            for f, kind in ACT_FIELDS.items():
                rr, cn, rn = compare(kind, match.get(f), ra.get(f))
                rows.append([cid, "E. Official actions", f"{label}: {f}", cn, match["_note"] if f.startswith("Basis") else "", rn, rr])
        for ca in coder_acts:
            if ca["_n"] not in used:
                rows.append([cid, "E. Official actions", f"Coder action {ca['_n']} (no reference match)", f"{s(ca['Action Type'])} ({s(ca['Category'])})", "", "", "Disagree (extra action)"])
        logs.append([cid, s(get("Coder|Minutes spent")[0]), s(get("Coder|Sources used")[0]), s(get("Coder|Uncertain points")[0])])

    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    fs = wb.create_sheet("Fields")
    fs.append(["Case ID", "Section", "Field", "Coder value", "Coder note", "Reference value", "Result", "Discussion / decision"])
    for r in rows:
        fs.append(r)
    n_f = fs.max_row
    bs = wb.create_sheet("Checkbox items")
    bs.append(["Case ID", "Group", "Item", "Coder (1 = marked)", "Reference (1 = marked)", "Agree"])
    for i, b in enumerate(binary, start=2):
        bs.append(b + [f"=IF(D{i}=E{i},1,0)"])
    n_b = bs.max_row
    lg = wb.create_sheet("Coder log")
    lg.append(["Case ID", "Minutes spent", "Sources used", "Points where the codebook did not say what to do"])
    for l in logs:
        lg.append(l)

    sm.append(["Agreement with the reference coding"])
    sm.append([f"Coder file: {coder_path}. Reference: {ref_path}. 'Partly' means the year matches or the sets overlap; 'Review' fields are free text and not scored."])
    sm.append([])
    sm.append(["Scored fields, by case", "Agree", "Partly", "Disagree", "Not coded", "Agreement"])
    for cid in coder:
        r = sm.max_row + 1
        sm.append([cid] + [f'=COUNTIFS(Fields!$A$2:$A${n_f},$A{r},Fields!$G$2:$G${n_f},"{x}*")' for x in ("Agree", "Partly", "Disagree", "Not coded")]
                  + [f"=IFERROR(B{r}/(B{r}+C{r}+D{r}),\"\")"])
    t = sm.max_row + 1
    sm.append(["All cases"] + [f"=SUM({c}5:{c}{t-1})" for c in "BCDE"] + [f"=IFERROR(B{t}/(B{t}+C{t}+D{t}),\"\")"])
    sm.append([])
    sm.append(["Checkbox items (Cohen's kappa)", "Items", "Observed agreement", "Coder marked", "Reference marked", "Expected agreement", "Kappa"])
    for grp in ("Category", "Factor", "Trigger", "All"):
        r = sm.max_row + 1
        cond = "" if grp == "All" else f",'Checkbox items'!$B$2:$B${n_b},\"{grp}\""
        rng = lambda col: f"'Checkbox items'!${col}$2:${col}${n_b}"
        if grp == "All":
            sm.append([grp, f"=COUNT({rng('F')})", f"=IFERROR(AVERAGE({rng('F')}),\"\")", f"=IFERROR(AVERAGE({rng('D')}),\"\")",
                       f"=IFERROR(AVERAGE({rng('E')}),\"\")"])
        else:
            sm.append([grp, f"=COUNTIFS({rng('B')},\"{grp}\")", f"=IFERROR(AVERAGEIFS({rng('F')}{cond}),\"\")",
                       f"=IFERROR(AVERAGEIFS({rng('D')}{cond}),\"\")", f"=IFERROR(AVERAGEIFS({rng('E')}{cond}),\"\")"])
        sm.cell(r, 6).value = f"=IFERROR(D{r}*E{r}+(1-D{r})*(1-E{r}),\"\")"
        sm.cell(r, 7).value = f"=IFERROR((C{r}-F{r})/(1-F{r}),\"\")"
    sm.append([])
    sm.append(["Kappa is undefined (shown blank) when both coders mark every item, or no item, the same way. With few cases, read it alongside the raw disagreements on the Fields sheet."])

    for ws, widths in ((sm, [34, 12, 18, 16, 18, 20, 12]), (fs, [15, 26, 46, 34, 44, 44, 22, 40]),
                       (bs, [15, 12, 56, 18, 20, 10]), (lg, [15, 14, 50, 80])):
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=F, size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for ws in (fs, bs, lg):
        for c in ws[1]:
            c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
        ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    sm["A1"].font = Font(name=F, bold=True, size=13)
    for row in sm.iter_rows():
        if row[0].value in ("Scored fields, by case", "Checkbox items (Cohen's kappa)"):
            for c in row:
                c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    for r in range(5, t + 1):
        sm.cell(r, 6).number_format = "0%"
    for r in range(t + 3, t + 7):
        for col in range(3, 8):
            sm.cell(r, col).number_format = "0.00"
    fs.conditional_formatting.add(f"A2:H{n_f}", FormulaRule(formula=[f'LEFT($G2,8)="Disagree"'], fill=PatternFill("solid", start_color="FBE3E3")))
    fs.conditional_formatting.add(f"A2:H{n_f}", FormulaRule(formula=[f'$G2="Partly"'], fill=PatternFill("solid", start_color="FFF4DC")))
    wb.save(out_path)
    print(f"{out_path}: {len(rows)} fields compared, {len(binary)} checkbox items")


if __name__ == "__main__":
    main(*sys.argv[1:4])
