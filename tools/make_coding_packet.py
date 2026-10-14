"""Build a blind coding template for an independent second coder.

Usage:
  python3 tools/make_coding_packet.py data/ARWC_Case_Inventory.xlsx coding/ARWC_Coding_Round1.xlsx \
      ARWC-TWN-001 ARWC-CHN-002 ARWC-CHN-001 ARWC-SGP-001 ARWC-IND-001

The template carries the case IDs, the names of the people convicted and the source list for
each case, and nothing of the reference coding. Fill it in, then run tools/compare_coding.py.
"""
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
HDR = PatternFill("solid", start_color="2F4A8A")
SEC = PatternFill("solid", start_color="E4E9F4")
INPUT = PatternFill("solid", start_color="FFF7D6")
thin = Side(style="thin", color="D0D7E4")

# Codebook v1.0 vocabularies (Supplement A)
YN = ["Yes", "No"]
YNU = ["Yes", "No", "Unknown"]
MARK = ["Marked", "Not marked"]
PRECISION = ["day", "month", "year"]
CV = {
    "gender": ["Male", "Female", "Other", "Unknown"],
    "sentence": ["Death (immediate)", "Death with reprieve", "Life imprisonment", "Fixed term", "Other (specify)"],
    "age_group": ["Minor", "Adult", "Mixed", "Unknown"],
    "relationship": ["Stranger", "Acquaintance", "Family member", "Other", "Unknown"],
    "finality": ["Ordinary appeals completed", "Ordinary appeal period expired", "Judgment otherwise entered into legal effect"],
    "action": ["Retrial acquittal", "Post-finality appellate acquittal", "Conviction vacated or annulled",
               "Extraordinary appeal decision quashing a final judgment", "Prosecutorial admission of error",
               "State compensation decision", "Administrative rectification", "Pardon (innocence-linked)",
               "Commutation (innocence-linked)", "Other (specify)"],
    "category": ["A", "B", "C", "None"],
    "basis": ["Insufficiency of the original evidence", "New evidence of innocence", "Procedural or legal error",
              "Other (specify)", "Not stated"],
    "passfail": ["PASS", "FAIL"],
    "inclusion": ["INCLUDED", "EXCLUDED", "PENDING"],
}
MULTI = {
    "offense": "Homicide; sexual offense; robbery; assault; kidnapping or abduction; arson; drug offense; fraud or property offense; public-order offense; other (specify). Several allowed, separated by semicolons.",
    "evidence": "Confession; eyewitness identification; forensic evidence (specify type); documentary evidence; informant or witness testimony; circumstantial evidence; other (specify). Several allowed, separated by semicolons.",
}
FACTORS = [("1", "False or coerced confession", "Confession evidence contributed to the conviction and was retracted, excluded or found unreliable."),
           ("1a", "Torture or physical coercion", "Where marked, also mark 4."),
           ("1b", "Prolonged detention or questioning without counsel", ""),
           ("1c", "Deception, inducement, or threats", ""),
           ("2", "Eyewitness misidentification", ""),
           ("2a", "Suggestive identification procedure", ""),
           ("2b", "Identification treated as decisive despite inconsistency or weak conditions", ""),
           ("3", "Forensic or evidentiary error", "Includes misapplied, overstated or invalid forensic evidence."),
           ("4", "Investigative or prosecutorial misconduct", ""),
           ("4a", "Suppression of exculpatory evidence", ""),
           ("4b", "Failure to act on exculpatory developments", "E.g. declining to pursue a confession by another person."),
           ("4c", "Fabrication of evidence", ""),
           ("5", "Inadequate defense", ""),
           ("6", "Perjury or false accusation", "Includes informant testimony."),
           ("7", "Systemic or structural issue", "Specify. Only where the documents show the feature was relevant at the time; never a label for a jurisdiction."),
           ("8", "Other contributing factor", "Specify.")]
TRIGGERS = [("Real perpetrator identified", "Another person identified as responsible, whether or not charged."),
            ("Confession by another person", "Recorded separately: a confession may not lead to identification or charge."),
            ("New or re-examined forensic evidence", "Includes re-testing of existing material."),
            ("Recantation of testimony", ""),
            ("Newly disclosed or discovered evidence", "Evidence not before the original court, including material previously withheld."),
            ("Sustained petitioning or family advocacy", ""),
            ("Involvement of an innocence organization or counsel", ""),
            ("Media or public attention", ""),
            ("Official review initiative", "Court-, procuratorate- or commission-initiated review; the impetus, not the route."),
            ("Other (specify)", "")]
N_ACTIONS = 4

# Sources from the website records, where the coded entry lists documents without numbers or links.
EXTRA_SOURCES = {
    "ARWC-CHN-001": [
        ["Judgment", "最高人民法院刑事判决书 (2016)最高法刑再3号 / Retrial criminal judgment", "Supreme People's Court", "2016-12-02", "https://www.court.gov.cn/zixun/xiangqing/33331.html"],
        ["Judgment", "(2016)最高法刑申188号 / Retrial decision", "Supreme People's Court", "2016-06-06", ""],
        ["Judgment", "(1995)冀刑一终字第129号 / Second-instance criminal judgment", "Hebei Higher People's Court", "1995-04-25", ""],
        ["Judgment", "(1995)石刑初字第53号 / First-instance criminal judgment", "Shijiazhuang Intermediate People's Court", "1995-03-15", ""],
        ["Compensation decision", "(2016)冀法赔1号 / State compensation decision", "Hebei Higher People's Court", "2017-03", ""],
        ["Media report", "聂树斌案获赔偿268万余元", "Xinhua", "2017-03-30", "http://www.xinhuanet.com/legal/2017-03/30/c_1120727450.htm"],
    ],
}


def fields_for(n_people):
    """(key, section, label, kind, guidance). Keys match the reference sheet's column names."""
    rows = [
        ("Case Name (Original Language)", "A. Case identification", "Case name (original language)", "text", "In the original script."),
        ("Original Language", "A. Case identification", "Original language", "text", "E.g. Chinese (Traditional)."),
    ]
    for p in range(1, n_people + 1):
        sec = "B. Individual" if n_people == 1 else f"B{p}. Individual"
        rows += [
            (f"B{p}|Individual Name (Original Language)", sec, "Name (original script)", "text", ""),
            (f"B{p}|Gender", sec, "Gender", "gender", ""),
            (f"B{p}|Age at Conviction", sec, "Age at conviction", "text", "Integer, years. Note which judgment you used."),
            (f"B{p}|Executed", sec, "Executed", "yn", ""),
            (f"B{p}|Date of Execution", sec, "Date of execution", "text", "YYYY-MM-DD, or leave blank."),
            (f"B{p}|Years Lost", sec, "Years lost", "text", "Decimal years, from detention to release or execution. Show your day count in the note."),
            (f"B{p}|Years Lost: Approximation Flag", sec, "Years lost: approximated?", "yn", ""),
            (f"B{p}|Posthumous Case", sec, "Posthumous case", "yn", "Corrective action after the person's death."),
        ]
    rows += [
        ("Offense Category", "C. Offense, sentence, victim", "Offense category", "offense", MULTI["offense"]),
        ("Offense as Charged (Original Language)", "C. Offense, sentence, victim", "Offense as charged (original language)", "text", ""),
        ("Offense as Charged (English Gloss)", "C. Offense, sentence, victim", "Offense as charged (English gloss)", "text", ""),
        ("Sentence", "C. Offense, sentence, victim", "Sentence", "sentence", "Add detail in the note."),
        ("Court of Original Trial", "C. Offense, sentence, victim", "Court of original trial", "text", "Name in the original language and English."),
        ("Evidence Relied Upon at Trial", "C. Offense, sentence, victim", "Evidence relied upon at trial", "evidence", MULTI["evidence"]),
        ("Number of Victims", "C. Offense, sentence, victim", "Number of victims", "text", "Integer."),
        ("Victim Age Group", "C. Offense, sentence, victim", "Victim age group", "age_group", ""),
        ("Victim–Defendant Relationship", "C. Offense, sentence, victim", "Victim–defendant relationship", "relationship", ""),
        ("Alleged Victim Later Reappeared", "C. Offense, sentence, victim", "Alleged victim later reappeared", "ynu", ""),
        ("Actual Perpetrator Identified", "C. Offense, sentence, victim", "Actual perpetrator identified", "ynu", ""),
        ("Subsequent Victims Attributed to Actual Perpetrator", "C. Offense, sentence, victim", "Subsequent victims attributed to actual perpetrator", "ynu", ""),
        ("Finality Basis", "D. Finality", "Finality basis", "finality", ""),
        ("Date of Finality", "D. Finality", "Date of finality", "text", "YYYY-MM-DD (or YYYY-MM, YYYY). Give the precision in the note."),
        ("Finality Note", "D. Finality", "Finality note", "text", "At most two sentences; case-specific facts only."),
        ("Category A", "E. Official action: categories", "Judicial reversal (A)", "mark", "A court ended, removed or nullified the finding of guilt after finality."),
        ("Category B", "E. Official action: categories", "Formal acknowledgment of error (B)", "mark", "A public authority's own instrument speaks to the validity of the conviction."),
        ("Category C", "E. Official action: categories", "Innocence-related relief (C)", "mark", "Relief tied to innocence concerns only by the surrounding record."),
    ]
    for a in range(1, N_ACTIONS + 1):
        sec = f"E. Official action {a}"
        hint = " Leave blank if there is no further action." if a > 1 else ""
        rows += [
            (f"Act{a}|Action Type", sec, "Action type", "action", "Controlled vocabulary." + hint),
            (f"Act{a}|Category", sec, "Category", "category", ""),
            (f"Act{a}|Document Type", sec, "Document type", "text", "Original language and English."),
            (f"Act{a}|Date of Action", sec, "Date of action", "text", "YYYY-MM-DD (or YYYY-MM, YYYY)."),
            (f"Act{a}|Issuing Authority", sec, "Issuing authority", "text", "Original language and English."),
            (f"Act{a}|Procedural Route", sec, "Procedural route", "text", "Local procedure, original-language name and English gloss."),
            (f"Act{a}|Basis Stated in the Decision", sec, "Basis stated in the decision", "basis", "What the decision itself says, not what prompted it."),
        ]
    rows += [
        ("Compensation: Compensation Awarded", "E. Compensation", "Compensation awarded", "ynu", ""),
        ("Compensation: Amount / Currency / Year", "E. Compensation", "Amount / currency / year", "text", "E.g. 4,960,521.50 / CNY / 2020."),
        ("Compensation: Decision Acknowledges Wrongful Conviction", "E. Compensation", "Decision acknowledges wrongful conviction", "yn", "If Yes, it is usually also recorded as a Category B action."),
    ]
    for code, name, hint in FACTORS:
        rows.append((f"{code} {name}: marked", "G. Contributing factors", f"{code}. {name}", "yn",
                     ("State the documentary basis in the note. " + hint).strip()))
    for name, hint in TRIGGERS:
        rows.append((f"{name}: marked", "H. Key triggers", name, "yn", hint))
    rows += [
        ("Criminal offence: status", "I. Inclusion verification", "Criminal offence", "passfail", ""),
        ("Finality: status", "I. Inclusion verification", "Finality", "passfail", ""),
        ("Official post-conviction action: status", "I. Inclusion verification", "Official post-conviction action", "passfail", ""),
        ("Documentation: status", "I. Inclusion verification", "Documentation", "passfail", "At least one traceable official document."),
        ("Inclusion determination: status", "I. Inclusion verification", "Inclusion determination", "inclusion", ""),
        ("Coder|Minutes spent", "Coder log", "Time spent on this case (minutes)", "text", "Include reading."),
        ("Coder|Sources used", "Coder log", "Sources you actually read", "text", "List any you added that are not on the Sources sheet."),
        ("Coder|Uncertain points", "Coder log", "Points where the codebook did not tell you what to do", "text", "The most useful field. One line per point."),
    ]
    return rows


def kind_values(kind):
    return {"yn": YN, "ynu": YNU, "mark": MARK}.get(kind) or CV.get(kind)


def main(inv_path, out_path, case_ids):
    inv = load_workbook(inv_path, read_only=True)
    coded = inv["Coded cases"]
    hdr = [c.value for c in next(coded.iter_rows(min_row=2, max_row=2))]
    ref = {r[hdr.index("Case ID")]: dict(zip(hdr, r)) for r in coded.iter_rows(min_row=3, values_only=True)}
    people = {}
    for r in inv["Individuals"].iter_rows(min_row=2, values_only=True):
        people.setdefault(r[0], []).append(r[3] or r[2])
    sources = {}
    for r in inv["Coded sources"].iter_rows(min_row=2, values_only=True):
        url = re.search(r"https?://\S+", r[5] or "")
        sources.setdefault(r[0], []).append([r[1], r[2], r[3], r[4], url.group(0).rstrip(";") if url else ""])
    missing = [c for c in case_ids if c not in ref]
    if missing:
        sys.exit(f"Not in the 'Coded cases' sheet: {missing}")

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    text = [
        ("ARWC independent coding, round 1", 14, True),
        ("Thank you for coding these cases. The aim is to test whether Codebook v1.0 leads two coders to the same coding, and to find where it does not.", 10, False),
        ("", 10, False),
        ("What to use", 11, True),
        ("1. Codebook v1.0: Supplement A of the framework paper (Zhang et al., 2026). The 'Codebook' sheet here summarises the fields.", 10, False),
        ("2. The ARWC website's Methodology and Glossary pages.", 10, False),
        ("3. The sources listed on the 'Sources' sheet, and any other official documents you find. Read original-language documents wherever you can.", 10, False),
        ("", 10, False),
        ("What not to use", 11, True),
        ("Please do not open the website's record pages for these cases, the coded-set documents or the inventory workbook until you have finished. They hold the reference coding.", 10, False),
        ("", 10, False),
        ("How to fill in a case sheet", 11, True),
        ("Yellow cells are for you. Column D takes the value: use the dropdown where there is one. Column E takes your note: for every factor or trigger you mark, say what document shows it.", 10, False),
        ("Mark a factor or trigger 'Yes' or 'No' for every row, so that a blank is never ambiguous.", 10, False),
        ("Where the codebook does not tell you what to do, choose what you think is right, and write the problem in 'Points where the codebook did not tell you what to do'. Those notes matter more than agreement.", 10, False),
        ("Dates: YYYY-MM-DD where the day is known; otherwise YYYY-MM or YYYY, with the precision in the note.", 10, False),
        ("Case summaries (Section F) are not needed for this round.", 10, False),
        ("", 10, False),
        ("When you are done", 11, True),
        ("Send the file back. It will be compared field by field with the reference coding (tools/compare_coding.py), and we will go through the disagreements together.", 10, False),
    ]
    for t, size, bold in text:
        ins.append([t])
        ins.cell(ins.max_row, 1).font = Font(name=F, size=size, bold=bold)
        ins.cell(ins.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ins.column_dimensions["A"].width = 110

    cb = wb.create_sheet("Codebook")
    cb.append(["Section", "Field", "Permitted values or guidance"])
    for c in cb[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    seen = set()
    for key, sec, label, kind, hint in fields_for(1):
        if re.match(r"E\. Official action [2-9]", sec) or sec == "Coder log":
            continue
        sec = re.sub(r"^E\. Official action 1$", "E. Official actions (each action)", sec)
        if (sec, label) in seen:
            continue
        seen.add((sec, label))
        vals = kind_values(kind)
        cb.append([sec, label, "; ".join(vals) + (f". {hint}" if hint else "") if vals else hint])
    for row in cb.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABC", [30, 44, 90]):
        cb.column_dimensions[col].width = w
    cb.freeze_panes = "A2"

    for cid in case_ids:
        r = ref[cid]
        names = people.get(cid, [""])
        ws = wb.create_sheet(cid)
        ws.append(["key", "Section", "Field", "Your value", "Note: documentary basis", "Permitted values / guidance"])
        for c in ws[1]:
            c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
        ws.append(["Case ID", "A. Case identification", "Case ID", cid, "", "Pre-filled"])
        ws.append(["Jurisdiction ID", "A. Case identification", "Jurisdiction ID", r.get("Jurisdiction ID", ""), "", "Pre-filled"])
        for i, n in enumerate(names, 1):
            label = "Person convicted" if len(names) == 1 else f"Person convicted (B{i})"
            ws.append([f"B{i}|Name", "B. Individual", label, n, "", "Pre-filled"])
        dvs = {}
        last_sec = None
        for key, sec, label, kind, hint in fields_for(len(names)):
            if sec != last_sec:
                ws.append(["", sec]); rr = ws.max_row
                for c in ws[rr]:
                    c.fill = SEC; c.font = Font(name=F, bold=True, size=10)
                last_sec = sec
            vals = kind_values(kind)
            ws.append([key, sec, label, "", "", ("; ".join(vals) + (". " if hint else "") if vals else "") + hint])
            rr = ws.max_row
            ws.cell(rr, 4).fill = INPUT; ws.cell(rr, 5).fill = INPUT
            if vals:
                if kind not in dvs:
                    dvs[kind] = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
                    ws.add_data_validation(dvs[kind])
                dvs[kind].add(f"D{rr}")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if not c.font.bold:
                    c.font = Font(name=F, size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(bottom=thin)
        for col, w in zip("ABCDEF", [8, 24, 40, 34, 54, 54]):
            ws.column_dimensions[col].width = w
        ws.column_dimensions["A"].hidden = True
        ws.freeze_panes = "D2"

    so = wb.create_sheet("Sources")
    so.append(["Case ID", "Source type", "Title (original / English)", "Authority / author", "Date", "Link"])
    for c in so[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    for cid in case_ids:
        for s in sources.get(cid, []) + EXTRA_SOURCES.get(cid, []):
            so.append([cid] + s)
            if s[4]:
                so.cell(so.max_row, 6).hyperlink = s[4]
    for row in so.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10, color="2F4A8A" if c.column == 6 and c.value else None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", [15, 20, 60, 32, 14, 60]):
        so.column_dimensions[col].width = w
    so.freeze_panes = "B2"

    wb.save(out_path)
    print(f"{out_path}: {len(case_ids)} cases, {sum(len(sources.get(c, [])) for c in case_ids)} sources")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3:])
