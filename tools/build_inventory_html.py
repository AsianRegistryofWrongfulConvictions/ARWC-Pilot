"""Build inventory.html from the case inventory workbook.

Usage: python3 tools/build_inventory_html.py data/ARWC_Case_Inventory.xlsx inventory.html

The workbook's "Cases" sheet is the source of truth: edit it, then rerun this script.
"""
import json
import sys
from datetime import date

from openpyxl import load_workbook

KEYS = ["inv", "arwc", "status", "person", "local", "jur", "place", "offense", "final", "corr",
        "outcome", "cat", "finality", "lead", "url", "notes"]


def read_rows(path):
    ws = load_workbook(path, read_only=True)["Cases"]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue
        row = {k: ("" if v is None else v) for k, v in zip(KEYS, values)}
        for k in ("final", "corr"):  # years stay numbers, or blank
            row[k] = int(row[k]) if str(row[k]).isdigit() else ""
        rows.append(row)
    return rows


PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>ARWC Case Inventory</title>
<meta name="description" content="Working inventory of ARWC records and publicly documented candidate cases for screening.">
<meta name="robots" content="noindex, nofollow">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600&family=Noto+Serif+SC:wght@400;600&family=Noto+Serif+TC:wght@400;600&family=Noto+Serif+JP:wght@400;600&family=Noto+Serif+KR:wght@400;600&display=swap">
<style>
:root{
  --paper:#F4F6F9; --surface:#FFFFFF; --ink:#172039; --ink-2:#4E5870;
  --indigo:#2F4A8A; --indigo-strong:#223A73; --wash:#E4E9F4;
  --gardenia:#B8841F; --amber-wash:#FBF1DC;
  --rule:#D0D7E4; --rule-strong:#9CA8C0;
  --sans:"Be Vietnam Pro","Segoe UI",system-ui,-apple-system,"Helvetica Neue",Arial,"PingFang SC","Hiragino Sans","Malgun Gothic",sans-serif;
  color-scheme:light;
}
@media (prefers-color-scheme: dark){
  :root:not([data-theme="light"]){
    --paper:#0E1322; --surface:#151C31; --ink:#E7EAF3; --ink-2:#A9B0C6;
    --indigo:#A9BBEE; --indigo-strong:#C6D2F6; --wash:#1B2543;
    --gardenia:#E2B458; --amber-wash:#2A2414;
    --rule:#28314E; --rule-strong:#44507A; color-scheme:dark;
  }
}
:root[data-theme="dark"]{
  --paper:#0E1322; --surface:#151C31; --ink:#E7EAF3; --ink-2:#A9B0C6;
  --indigo:#A9BBEE; --indigo-strong:#C6D2F6; --wash:#1B2543;
  --gardenia:#E2B458; --amber-wash:#2A2414;
  --rule:#28314E; --rule-strong:#44507A; color-scheme:dark;
}
*,*::before,*::after{box-sizing:border-box}
body{margin:0;background:var(--paper);color:var(--ink);font-family:var(--sans);font-size:1rem;line-height:1.55}
:lang(zh-Hans){font-family:"Noto Serif SC","Songti SC",serif}
:lang(zh-Hant){font-family:"Noto Serif TC","Songti TC",serif}
:lang(ja){font-family:"Noto Serif JP","Hiragino Mincho ProN",serif}
:lang(ko){font-family:"Noto Serif KR","AppleMyungjo",serif}
a{color:var(--indigo);text-underline-offset:.18em}
:focus-visible{outline:3px solid var(--gardenia);outline-offset:2px}
.wrap{max-width:1240px;margin:0 auto;padding-inline:clamp(1rem,4vw,2.5rem)}
.bar{background:var(--amber-wash);border-bottom:1px solid var(--rule);font-size:.86rem;padding:.55rem 0}
header.top{border-bottom:1px solid var(--rule);padding:.9rem 0}
header.top .wrap{display:flex;justify-content:space-between;align-items:center;gap:1rem;flex-wrap:wrap}
.brand{font-weight:600;letter-spacing:.03em;color:var(--ink);text-decoration:none}
.brand span{font-weight:400;color:var(--ink-2);margin-left:.5rem;letter-spacing:0}
main{padding:2.2rem 0 4rem}
h1{font-size:clamp(1.9rem,2.5vw + 1rem,2.6rem);font-weight:500;letter-spacing:-.02em;margin:0}
.lead{color:var(--ink-2);max-width:72ch;margin:.8rem 0 0;font-size:1.05rem}
.actions{display:flex;gap:.6rem;flex-wrap:wrap;margin-top:1.2rem}
.btn{display:inline-flex;align-items:center;gap:.4rem;font:inherit;font-size:.93rem;font-weight:500;padding:.55rem .95rem;border-radius:6px;border:1.5px solid var(--indigo);color:var(--indigo);background:transparent;text-decoration:none}
.btn.primary{background:var(--indigo);color:var(--paper)}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(9.5rem,1fr));gap:.7rem;margin:2rem 0 0}
.tile{background:var(--surface);border:1px solid var(--rule);border-radius:8px;padding:.75rem .9rem;text-align:left;font:inherit;color:inherit;cursor:pointer}
.tile[aria-pressed="true"]{border-color:var(--indigo);box-shadow:inset 0 0 0 1px var(--indigo)}
.tile .n{display:block;font-size:1.6rem;font-weight:600;font-variant-numeric:tabular-nums;line-height:1.1}
.tile .l{display:block;font-size:.82rem;color:var(--ink-2);margin-top:.2rem}
.controls{display:grid;grid-template-columns:2fr repeat(3,1fr);gap:.7rem;align-items:end;margin:1.6rem 0 .8rem}
.field label{display:block;font-size:.82rem;color:var(--ink-2);margin-bottom:.25rem}
.field input,.field select{width:100%;font:inherit;font-size:.93rem;color:var(--ink);background:var(--surface);border:1px solid var(--rule-strong);border-radius:6px;padding:.5rem .6rem;min-height:2.5rem}
@media (max-width:820px){.controls{grid-template-columns:1fr 1fr}.controls .search{grid-column:1/-1}}
.result{display:flex;justify-content:space-between;gap:1rem;flex-wrap:wrap;font-size:.9rem;color:var(--ink-2);margin-bottom:.6rem}
.linklike{background:none;border:0;padding:0;font:inherit;color:var(--indigo);text-decoration:underline;cursor:pointer}
.scroll{overflow-x:auto;border:1px solid var(--rule);border-radius:10px;background:var(--surface)}
table{border-collapse:collapse;width:100%;font-size:.88rem}
th,td{text-align:left;padding:.6rem .7rem;border-bottom:1px solid var(--rule);vertical-align:top}
thead th{position:sticky;top:0;background:var(--paper);font-size:.78rem;font-weight:600;color:var(--ink-2);white-space:nowrap;z-index:1}
thead th button{all:unset;cursor:pointer}
thead th button::after{content:" ↕";color:var(--rule-strong)}
thead th[aria-sort="ascending"] button::after{content:" ↑";color:var(--indigo)}
thead th[aria-sort="descending"] button::after{content:" ↓";color:var(--indigo)}
tbody tr:last-child td{border-bottom:0}
td.num{font-variant-numeric:tabular-nums;white-space:nowrap}
.who{font-weight:500}
.local{display:block;color:var(--ink-2);font-size:.86rem}
.muted{color:var(--ink-2)}
.pill{display:inline-block;font-size:.74rem;font-weight:600;padding:.08rem .5rem;border-radius:99px;white-space:nowrap;border:1px solid var(--rule-strong);color:var(--ink-2)}
.pill.v{background:var(--indigo);border-color:var(--indigo);color:var(--paper)}
.pill.p{background:var(--wash);border-color:var(--indigo);color:var(--indigo-strong)}
.pill.pend{background:var(--amber-wash);border-color:var(--gardenia);color:var(--ink)}
.pill.x{border-style:dashed}
.fin{font-size:.8rem;white-space:nowrap}
.fin.Likely,.fin.Pass{color:var(--indigo-strong)}
.fin.Doubtful,.fin.Fail{color:var(--ink-2);text-decoration:underline dotted}
.empty{padding:2rem 1rem}
footer{border-top:1px solid var(--rule);padding:1.8rem 0 2.4rem;font-size:.85rem;color:var(--ink-2)}
@media (max-width:760px){
  thead{position:absolute;width:1px;height:1px;overflow:hidden;clip:rect(0 0 0 0)}
  tbody tr{display:block;padding:.8rem .9rem;border-bottom:1px solid var(--rule)}
  td{display:grid;grid-template-columns:7rem 1fr;gap:.5rem;border:0;padding:.15rem 0}
  td::before{content:attr(data-label);color:var(--ink-2);font-size:.8rem}
  td:empty{display:none}
}
</style>
</head>
<body>
<div class="bar"><div class="wrap"><strong>Internal working list.</strong> Candidate rows are leads drawn from existing archives and public reporting. They have not been screened against the ARWC criteria and are not registry records.</div></div>
<header class="top"><div class="wrap"><a class="brand" href="index.html">ARWC<span>Asian Registry of Wrongful Convictions</span></a><a href="index.html#/registry">Published registry</a></div></header>
<main><div class="wrap">
  <h1>Case inventory</h1>
  <p class="lead">Every case already in the registry, together with publicly documented cases from the China Wrongful Convictions Case Archive, innocence projects and public reporting that could be screened for inclusion. Use the finality check to prioritise: ARWC includes only convictions corrected after they became final.</p>
  <div class="actions">
    <a class="btn primary" href="data/ARWC_Case_Inventory.xlsx" download>Download the Excel workbook</a>
  </div>
  <div class="tiles" id="tiles" role="group" aria-label="Filter by status"></div>
  <div class="controls">
    <div class="field search"><label for="q">Search names, places, offenses or notes</label><input id="q" type="search" autocomplete="off"></div>
    <div class="field"><label for="fJur">Jurisdiction</label><select id="fJur"></select></div>
    <div class="field"><label for="fFin">Finality check</label><select id="fFin"></select></div>
    <div class="field"><label for="fLead">Lead source</label><select id="fLead"></select></div>
  </div>
  <div class="result"><span id="count" aria-live="polite"></span><button type="button" class="linklike" id="clear">Clear filters</button></div>
  <div class="scroll">
    <table>
      <caption class="muted" style="position:absolute;left:-9999px">Case inventory</caption>
      <thead><tr>
        <th scope="col" data-k="status"><button type="button">Status</button></th>
        <th scope="col" data-k="person"><button type="button">Person(s)</button></th>
        <th scope="col" data-k="jur"><button type="button">Jurisdiction</button></th>
        <th scope="col">Offense</th>
        <th scope="col" data-k="final"><button type="button">Final</button></th>
        <th scope="col" data-k="corr"><button type="button">Corrected</button></th>
        <th scope="col">Correction, as reported</th>
        <th scope="col" data-k="finality"><button type="button">Finality</button></th>
        <th scope="col">Lead source</th>
        <th scope="col">Notes</th>
      </tr></thead>
      <tbody id="rows"></tbody>
    </table>
  </div>
</div></main>
<footer><div class="wrap">Generated __DATE__ from data/ARWC_Case_Inventory.xlsx by tools/build_inventory_html.py. Edit the workbook, then rerun the script. Years and outcomes are as reported by the lead source and have not been checked against judgments.</div></footer>
<script>
const ROWS = __DATA__;
const STATUS_ORDER = ["Verified record","Provisional record","Pending","Candidate - screened in","Candidate - not screened","Candidate - screened out","Set aside in pilot","Excluded (worked example)"];
const PILL = { "Verified record":"v", "Provisional record":"p", "Pending":"pend", "Set aside in pilot":"x", "Excluded (worked example)":"x", "Candidate - screened out":"x" };
const SHORT = { "Candidate - not screened":"Candidate", "Excluded (worked example)":"Excluded", "Set aside in pilot":"Set aside" };
const esc = s => String(s ?? "").replace(/[&<>"']/g, c => ({ "&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;" }[c]));
const LANG = { China:"zh-Hans", Taiwan:"zh-Hant", Japan:"ja", Korea:"ko", Israel:"he", Pakistan:"ur", Vietnam:"vi" };
const $ = id => document.getElementById(id);
const uniq = k => [...new Set(ROWS.map(r => r[k]).filter(Boolean))].sort();
const opts = (el, vals) => el.innerHTML = `<option value="">All</option>` + vals.map(v => `<option>${esc(v)}</option>`).join("");
opts($("fJur"), uniq("jur")); opts($("fFin"), uniq("finality")); opts($("fLead"), uniq("lead"));
let status = "", sortKey = "status", sortDir = 1;

const statuses = STATUS_ORDER.filter(s => ROWS.some(r => r.status === s)).concat(uniq("status").filter(s => !STATUS_ORDER.includes(s)));
function tiles() {
  $("tiles").innerHTML = [["", "All cases", ROWS.length], ...statuses.map(s => [s, s, ROWS.filter(r => r.status === s).length])]
    .map(([v, l, n]) => `<button type="button" class="tile" data-s="${esc(v)}" aria-pressed="${status === v}"><span class="n">${n}</span><span class="l">${esc(l)}</span></button>`).join("");
}
$("tiles").addEventListener("click", e => { const b = e.target.closest(".tile"); if (!b) return; status = b.dataset.s; render(); });

function cmp(a, b) {
  if (sortKey === "status") return (STATUS_ORDER.indexOf(a.status) - STATUS_ORDER.indexOf(b.status)) * sortDir || String(a.jur).localeCompare(b.jur);
  const x = a[sortKey], y = b[sortKey];
  if (x === "" && y !== "") return 1; if (y === "" && x !== "") return -1;
  return (typeof x === "number" ? x - y : String(x).localeCompare(String(y))) * sortDir;
}

function render() {
  tiles();
  const q = $("q").value.trim().toLowerCase();
  const list = ROWS.filter(r => (!status || r.status === status)
    && (!$("fJur").value || r.jur === $("fJur").value)
    && (!$("fFin").value || r.finality === $("fFin").value)
    && (!$("fLead").value || r.lead === $("fLead").value)
    && (!q || Object.values(r).join(" ").toLowerCase().includes(q))).sort(cmp);
  $("count").textContent = `Showing ${list.length} of ${ROWS.length} cases`;
  $("clear").hidden = !(status || q || $("fJur").value || $("fFin").value || $("fLead").value);
  document.querySelectorAll("thead th[data-k]").forEach(th => th.setAttribute("aria-sort", th.dataset.k === sortKey ? (sortDir > 0 ? "ascending" : "descending") : "none"));
  $("rows").innerHTML = list.length ? list.map(r => {
    const name = r.arwc && /^ARWC-/.test(r.arwc) && r.status !== "Pending" ? `<a href="index.html#/record/${esc(r.arwc)}">${esc(r.person)}</a>` : esc(r.person);
    const src = r.url ? `<a href="${esc(r.url)}" rel="noopener">${esc(r.lead)}</a>` : esc(r.lead);
    const loc = r.local ? `<span class="local" lang="${LANG[r.jur] || "en"}"${["Israel","Pakistan"].includes(r.jur) ? ' dir="rtl"' : ""}>${esc(r.local)}</span>` : "";
    return `<tr>
      <td data-label="Status"><span class="pill ${PILL[r.status] || ""}">${esc(SHORT[r.status] || r.status)}</span>${r.arwc ? `<span class="local num">${esc(r.arwc)}</span>` : ""}</td>
      <td data-label="Person(s)"><span class="who">${name}</span>${loc}</td>
      <td data-label="Jurisdiction">${esc(r.jur)}${r.place ? `<span class="local">${esc(r.place)}</span>` : ""}</td>
      <td data-label="Offense">${esc(r.offense)}</td>
      <td data-label="Final" class="num">${esc(r.final)}</td>
      <td data-label="Corrected" class="num">${esc(r.corr)}</td>
      <td data-label="Correction">${esc(r.outcome)}${r.cat ? ` <span class="muted">(${esc(r.cat)})</span>` : ""}</td>
      <td data-label="Finality"><span class="fin ${esc(String(r.finality).split(" ")[0])}">${esc(r.finality)}</span></td>
      <td data-label="Lead source">${src}</td>
      <td data-label="Notes" class="muted">${esc(r.notes)}</td>
    </tr>`;
  }).join("") : `<tr><td colspan="10" class="empty">No cases match these filters.</td></tr>`;
}
["q","fJur","fFin","fLead"].forEach(id => $(id).addEventListener("input", render));
$("clear").addEventListener("click", () => { status = ""; ["q","fJur","fFin","fLead"].forEach(id => $(id).value = ""); render(); });
document.querySelectorAll("thead th[data-k] button").forEach(b => b.addEventListener("click", () => {
  const k = b.parentElement.dataset.k; sortDir = sortKey === k ? -sortDir : 1; sortKey = k; render();
}));
render();
</script>
</body>
</html>
"""


def main(src, out):
    rows = read_rows(src)
    data = json.dumps(rows, ensure_ascii=False).replace("</", "<\\/")
    html = PAGE.replace("__DATA__", data).replace("__DATE__", date.today().strftime("%d %B %Y").lstrip("0"))
    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"{out}: {len(rows)} cases")


if __name__ == "__main__":
    main(*sys.argv[1:3])
